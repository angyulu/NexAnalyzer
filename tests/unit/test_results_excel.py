"""Unit tests for modules.spectra.io.results_excel (sample-results workbook)."""

from io import BytesIO

import numpy as np
import pytest
from openpyxl import load_workbook

from core.report.models import RAW_STAT_LABEL, PeakStat
from modules.spectra.io.results_excel import TechniqueResults, build_sample_results_xlsx
from modules.spectra.models.peak import FittedPeak, FitResult
from modules.spectra.models.spectrum import ProcessingSettings, SpectrumData, SpectrumFile
from modules.spectra.processing.peak_metrics import peak_intensity


def _fitted_peak(label="Si", center=520.0, area=1000.0, fwhm=8.0, x=None):
    component = np.exp(-0.5 * ((x - center) / (fwhm / 2.355)) ** 2) * 137.0
    return FittedPeak(
        label=label, center=center, center_stderr=0.25,
        area=area, area_stderr=10.0,
        width_fwhm=fwhm, width_stderr=0.5, shape=0.5,
        component_curve=component, color="#1f77b4",
    )


def _spectrum(mode="Raman", peaks=("Si",), flat=False):
    x = np.linspace(400.0, 700.0, 200)
    y = np.zeros_like(x) if flat else 50.0 * np.exp(-0.5 * ((x - 550.0) / 20.0) ** 2)
    data = SpectrumData(X=x, Y=y)
    fitted = [_fitted_peak(label=label, x=x) for label in peaks]

    return SpectrumFile(
        filename=f"{mode}_1", mode=mode,
        original_data=data, raw_data=data, processed_data=data,
        processing_settings=ProcessingSettings(),
        fit_result=FitResult(
            success=True, fitted_peaks=fitted,
            total_fit_curve=y, residuals=np.zeros_like(y),
            chi_squared=1.5, r_squared=0.987, convergence_time=0.25,
        ),
    )


def _peak_stat(label="Si", fwhm_mean=8.0, fwhm_std=0.4):
    return PeakStat(
        label=label, n=3,
        center_mean=520.1, center_std=0.3,
        intensity_mean=137.0, intensity_std=4.0,
        fwhm_mean=fwhm_mean, fwhm_std=fwhm_std,
    )


def _build(**kwargs):
    defaults = dict(
        sample_name="HADG06", material_name="WSe2", report_date="2026-08-24",
        techniques=[],
    )
    defaults.update(kwargs)
    return load_workbook(BytesIO(build_sample_results_xlsx(**defaults)))


def _rows(ws):
    """Sheet rows as lists of values, header included."""
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _column(ws, header):
    """One column's data values, keyed by its header text."""
    rows = _rows(ws)
    index = rows[0].index(header)
    return [row[index] for row in rows[1:]]


def _cells(ws):
    """Every non-empty cell's value, flattened — for asserting on the Summary
    sheet, whose blocks sit at row offsets no test should hardcode."""
    return [value for row in _rows(ws) for value in row if value is not None]


class TestSheets:
    def test_summary_comes_first_so_it_is_what_opens(self):
        wb = _build(techniques=[
            TechniqueResults("Raman", [(1, _spectrum())], [_peak_stat()]),
        ])

        assert wb.sheetnames[0] == "Summary"
        assert wb.sheetnames == ["Summary", "Raman"]

    def test_a_technique_with_no_fits_gets_no_sheet(self):
        """An empty sheet named "PL" claims a technique that was never
        measured."""
        wb = _build(techniques=[
            TechniqueResults("Raman", [(1, _spectrum())], [_peak_stat()]),
            TechniqueResults("PL", [], None, include_raw_row=True),
        ])

        assert "PL" not in wb.sheetnames

    def test_a_sample_with_nothing_fitted_still_carries_its_identity(self):
        wb = _build(techniques=[TechniqueResults("Raman", [], None, errors=[(1, "no peaks")])])

        assert wb.sheetnames == ["Summary"]
        assert "HADG06" in _cells(wb["Summary"])


class TestPointSheets:
    def test_one_row_per_peak_per_point(self):
        spectra = [(point, _spectrum(peaks=("LA", "E2g+A1g"))) for point in (1, 2, 3)]
        ws = _build(techniques=[TechniqueResults("Raman", spectra, [_peak_stat()])])["Raman"]

        assert _column(ws, "Point") == [1, 1, 2, 2, 3, 3]
        assert _column(ws, "Peak_Label") == ["LA", "E2g+A1g"] * 3

    def test_points_are_ordered_however_they_arrive(self):
        spectra = [(7, _spectrum()), (2, _spectrum()), (9, _spectrum())]
        ws = _build(techniques=[TechniqueResults("Raman", spectra, None)])["Raman"]

        assert _column(ws, "Point") == [2, 7, 9]

    def test_intensity_is_the_curve_maximum_not_the_area(self):
        """The bug that shipped a .pptx disagreeing with its own CSV. Area and
        intensity differ by ~FWHM x 1.064, so an area here would be silently
        wrong under a column headed Intensity."""
        spectrum = _spectrum()
        peak = spectrum.fit_result.fitted_peaks[0]
        ws = _build(techniques=[TechniqueResults("Raman", [(1, spectrum)], None)])["Raman"]

        assert _column(ws, "Intensity") == [pytest.approx(peak_intensity(peak))]
        assert _column(ws, "Intensity")[0] != pytest.approx(peak.area)

    def test_values_are_written_unrounded(self):
        """The number formats decide what Excel shows; the cell keeps the
        number, or a user widening the format gets digits we threw away."""
        spectrum = _spectrum()
        spectrum.fit_result.fitted_peaks[0].center = 520.123456789
        ws = _build(techniques=[TechniqueResults("Raman", [(1, spectrum)], None)])["Raman"]

        assert _column(ws, "Center") == [520.123456789]

    def test_the_source_file_is_named_so_a_row_traces_back_to_disk(self):
        ws = _build(techniques=[TechniqueResults(
            "Raman", [(1, _spectrum())], None,
            source_files={1: r"C:\samples\HADG06\RM_1.txt"},
        )])["Raman"]

        assert _column(ws, "Source_File") == ["RM_1.txt"]

    def test_fit_quality_rides_along_on_every_peak_row(self):
        ws = _build(techniques=[TechniqueResults("Raman", [(1, _spectrum())], None)])["Raman"]

        assert _column(ws, "R_Squared") == [0.987]
        assert _column(ws, "Chi_Squared") == [1.5]
        assert _column(ws, "Convergence_Time_s") == [0.25]


class TestEmpiricalRow:
    def test_pl_leads_each_point_with_its_raw_measurement(self):
        spectra = [(point, _spectrum(mode="PL")) for point in (1, 2)]
        ws = _build(techniques=[TechniqueResults(
            "PL", spectra, None, include_raw_row=True
        )])["PL"]

        assert _column(ws, "Peak_Label") == [RAW_STAT_LABEL, "Si", RAW_STAT_LABEL, "Si"]

    def test_raman_does_not_get_one(self):
        """Which techniques carry an empirical row is a reporting decision, and
        the master CSV draws the same line: PL only."""
        ws = _build(techniques=[TechniqueResults("Raman", [(1, _spectrum())], None)])["Raman"]

        assert RAW_STAT_LABEL not in _column(ws, "Peak_Label")

    def test_the_raw_row_carries_no_fit_columns(self):
        ws = _build(techniques=[TechniqueResults(
            "PL", [(1, _spectrum(mode="PL"))], None, include_raw_row=True
        )])["PL"]
        raw = _rows(ws)[1]
        headers = _rows(ws)[0]

        for header in ("Center_Stderr", "Intensity_Stderr", "R_Squared", "Shape"):
            assert raw[headers.index(header)] is None, f"{header} is not a measurement of a Raw row"

    def test_a_flat_spectrum_reports_its_intensity_but_withholds_a_width(self):
        """0 counts is a measurement and is reported. A width is not: a flat
        signal has no half-maximum crossing, so the cell stays empty rather
        than printing a 0.0 that would read as a measured width. Same split
        the master CSV and the report's Raw row make."""
        ws = _build(techniques=[TechniqueResults(
            "PL", [(1, _spectrum(mode="PL", flat=True))], None, include_raw_row=True
        )])["PL"]

        assert _column(ws, "Peak_Label") == [RAW_STAT_LABEL, "Si"]
        assert _column(ws, "Intensity")[0] == 0.0
        assert _column(ws, "FWHM")[0] is None


class TestSummarySheet:
    def test_mean_and_std_are_separate_numeric_columns(self):
        """Not one "520.1 ± 0.3" string: the point of a workbook is that the
        next person can compute with it."""
        ws = _build(techniques=[
            TechniqueResults("Raman", [(1, _spectrum())], [_peak_stat()]),
        ])["Summary"]
        values = _cells(ws)

        assert 520.1 in values and 0.3 in values
        assert not any(isinstance(v, str) and "±" in v and "520" in v for v in values)

    def test_an_unmeasurable_width_is_left_empty_not_zero(self):
        """The cell the .pptx dashes out. A 0.0 here would read as a width the
        sample was measured to have."""
        raw_stat = _peak_stat(label=RAW_STAT_LABEL, fwhm_mean=None, fwhm_std=None)
        ws = _build(techniques=[TechniqueResults(
            "PL", [(1, _spectrum(mode="PL"))], [raw_stat], include_raw_row=True
        )])["Summary"]
        rows = _rows(ws)
        header = next(row for row in rows if row and row[0] == "Peak")
        stat_row = rows[rows.index(header) + 1]

        assert stat_row[header.index("FWHM_Mean")] is None
        assert stat_row[header.index("FWHM_Std")] is None

    def test_a_table_led_by_the_empirical_row_is_not_called_a_fit_summary(self):
        wb = _build(techniques=[
            TechniqueResults("Raman", [(1, _spectrum())], [_peak_stat()]),
            TechniqueResults(
                "PL", [(1, _spectrum(mode="PL"))],
                [_peak_stat(label=RAW_STAT_LABEL), _peak_stat()],
                include_raw_row=True,
            ),
        ])
        captions = [v for v in _cells(wb["Summary"]) if isinstance(v, str) and "summary" in v]

        assert "Raman fit summary (mean ± std)" in captions
        assert "PL summary (empirical + fitted, mean ± std)" in captions

    def test_ratios_are_written_as_median_and_mad(self):
        ws = _build(techniques=[TechniqueResults(
            "Raman", [(1, _spectrum())], [_peak_stat()],
            ratios=[("LA / E2g+A1g", (0.123, 0.004, 9))],
        )])["Summary"]
        values = _cells(ws)

        assert "LA / E2g+A1g" in values
        assert 0.123 in values and 0.004 in values

    def test_excluded_points_are_named_not_just_absent(self):
        """A workbook that drops points silently reads as a complete record of
        a sample that was never measured that way."""
        ws = _build(techniques=[
            TechniqueResults("Raman", [(1, _spectrum())], [_peak_stat()], errors=[(4, "no peaks found")]),
            TechniqueResults("PL", [], None, errors=[(4, "fit did not converge")]),
        ])["Summary"]
        values = _cells(ws)

        assert any(isinstance(v, str) and "Excluded" in v for v in values)
        assert "no peaks found" in values
        assert "fit did not converge" in values

    def test_the_identity_block_names_sample_material_and_date(self):
        ws = _build(techniques=[TechniqueResults("Raman", [(1, _spectrum())], None)])["Summary"]
        values = _cells(ws)

        assert "HADG06" in values
        assert "WSe2" in values
        assert "2026-08-24" in values


class TestUsability:
    def test_point_sheets_freeze_the_header_and_offer_a_filter(self):
        spectra = [(point, _spectrum()) for point in range(1, 10)]
        ws = _build(techniques=[TechniqueResults("Raman", spectra, None)])["Raman"]

        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None
