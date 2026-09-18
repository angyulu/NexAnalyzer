"""
Integration test for the Plot Explorer page (pages/4_Plot_Explorer.py).

Drives the real Streamlit page through streamlit.testing.v1.AppTest, which is
the only way to cover the page's own wiring: the guards that stop the script
before a workbook is chosen, the session-state round-trip, and the fact that a
bad column choice surfaces as an st.error rather than a traceback.

The workbook is seeded into session state rather than picked, because the page
opens a native OS dialog and there is no dialog in a test process. The fixture
carries the awkward shapes the page exists to handle — two header rows,
duplicate column names, a spacer column, mixed numeric/text columns and blank
filler rows.
"""

import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from core.paths import PROJECT_ROOT
from modules.dataviz.io import config_store
from modules.dataviz.viz.scatter import DEFAULT_CONFIG

PAGE = str(PROJECT_ROOT / "pages" / "4_Plot_Explorer.py")


@pytest.fixture
def workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "HA_SplitTable"
    ws.append(["WaferID", "T", "H2Se", "H2O", "H2O", None, "FWHM", "P"])
    ws.append([None, "degC", "sccm", "sccm", "torr", None, "E2g", "torr"])
    ws.append(["HAAA01", 890, 2, 50, 175, None, 3.1, 65])
    ws.append(["HAAA02", 950, "1->4(0.2sccm)", 65, 150, None, 3.4, 65])
    ws.append([None, None, None, None, None, None, None, None])
    ws.append(["HAAA03", 700, 3, 50, 175, None, 3.9, 65])
    path = tmp_path / "schedule.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    """Keep the page's remembered settings out of the repo's data/ directory."""
    monkeypatch.setattr(config_store, "get_config_path", lambda: tmp_path / "plot_explorer.json")


def _seed(workbook, **overrides):
    """Session state as if a workbook and sheet had already been picked."""
    state = {
        "path": workbook,
        "sheet": "HA_SplitTable",
        "header_rows": 2,
        "coercions": {},
        "filters": [],
        "config": dict(DEFAULT_CONFIG),
        "sheets": None,
        "load_error": None,
    }
    state.update(overrides)
    return state


def _run(workbook, **overrides):
    at = AppTest.from_file(PAGE, default_timeout=90)
    at.session_state["plot_explorer"] = _seed(workbook, **overrides)
    return at.run()


class TestGuards:
    def test_renders_with_no_workbook_chosen(self):
        at = AppTest.from_file(PAGE, default_timeout=90).run()
        assert not at.exception
        # Stops at the first section rather than rendering empty controls.
        assert any("Select an .xlsx" in i.value for i in at.info)

    def test_a_workbook_that_has_moved_is_an_error_not_a_traceback(self, tmp_path):
        at = AppTest.from_file(PAGE, default_timeout=90)
        at.session_state["plot_explorer"] = _seed(str(tmp_path / "gone.xlsx"))
        at.run()
        assert not at.exception
        assert any("no longer at" in e.value for e in at.error)


class TestLoading:
    def test_sheet_loads_and_columns_are_joined(self, workbook):
        at = _run(workbook)
        assert not at.exception
        # Two header rows joined, the duplicate suffixed, the spacer dropped,
        # and the blank filler row gone: 3 data rows, 7 of the sheet's 8
        # columns.
        assert any("3 rows" in c.value and "7 columns" in c.value for c in at.caption)

    def test_the_column_dropdowns_offer_the_joined_names(self, workbook):
        """The names the rest of the page is driven by, asserted once.

        Two header rows joined ("T (degC)"), the repeated H2O separated by its
        qualifier, and the headerless spacer column absent.
        """
        at = _run(workbook)
        x_dropdown = [s for s in at.selectbox if s.label == "X"][0]
        # options[0] is the "none" placeholder.
        assert x_dropdown.options[1:] == [
            "WaferID", "T (degC)", "H2Se (sccm)", "H2O (sccm)", "H2O (torr)",
            "FWHM (E2g)", "P (torr)",
        ]

    def test_mixed_column_is_announced(self, workbook):
        at = _run(workbook)
        assert any("hold both numbers and text" in c.value for c in at.caption)

    def test_prompts_for_axes_before_plotting(self, workbook):
        at = _run(workbook)
        assert any("Choose an X and a Y" in i.value for i in at.info)


class TestPlotting:
    def test_a_configured_scatter_renders(self, workbook):
        at = _run(workbook, config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)"})
        assert not at.exception
        assert not at.error
        assert len(at.get("plotly_chart")) == 1

    def test_coercion_makes_a_text_column_plottable(self, workbook):
        at = _run(
            workbook,
            coercions={"H2Se (sccm)": "leading"},
            config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "H2Se (sccm)"},
        )
        assert not at.exception
        assert len(at.get("plotly_chart")) == 1

    def test_a_size_on_a_text_column_is_unreachable_through_the_page(self, workbook):
        """The Size dropdown offers numeric columns only, so a stale remembered
        `size` naming a text column is reset to unset rather than reaching
        `build_scatter` -- which would raise on it, as the unit tests assert.
        The page corrects it silently and still draws.
        """
        at = _run(workbook, config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)",
                                    "size": "WaferID"})
        assert not at.exception
        assert not at.error
        assert len(at.get("plotly_chart")) == 1
        assert at.session_state["plot_explorer"]["config"]["size"] is None

    def test_a_numeric_range_filter_narrows_the_plot(self, workbook):
        at = _run(
            workbook,
            filters=[{"column": "T (degC)", "minimum": 800.0, "maximum": 1000.0}],
            config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)"},
        )
        assert not at.exception
        assert any("2 of 3 rows kept" in c.value for c in at.caption)

    def test_a_constant_numeric_column_does_not_break_the_slider(self, workbook):
        """st.slider raises when its bounds are equal, which a parameter held
        constant across every run produces."""
        at = _run(
            workbook,
            filters=[{"column": "P (torr)", "minimum": 65.0, "maximum": 65.0}],
            config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)"},
        )
        assert not at.exception
        assert any("nothing to narrow" in c.value for c in at.caption)

    def test_a_filter_that_excludes_everything_warns(self, workbook):
        at = _run(
            workbook,
            filters=[{"column": "WaferID", "keep": ["nonexistent-wafer"]}],
            config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)"},
        )
        assert not at.exception
        assert any("No rows survive" in w.value for w in at.warning)


@pytest.fixture
def two_sheet_book(tmp_path):
    """Three sheets: two with different columns, two with IDENTICAL columns.

    The identical pair mirrors `VAHA_SplitTable` and `VB_SplitTable` in the real
    process workbook, whose column sets are byte-identical.
    """
    wb = Workbook()
    a = wb.active
    a.title = "SheetA"
    a.append(["Alpha", "Beta"])
    a.append([1, 2])
    a.append([3, 4])

    b = wb.create_sheet("SheetB")
    b.append(["Gamma", "Delta", "Epsilon"])
    b.append([5, 6, 7])
    b.append([8, 9, 10])

    twin = wb.create_sheet("SheetATwin")
    twin.append(["Alpha", "Beta"])
    twin.append([100, 200])
    twin.append([300, 400])

    path = tmp_path / "sheets.xlsx"
    wb.save(path)
    return str(path)


def _sheet_box(at):
    return [s for s in at.selectbox if s.label == "Sheet"][0]


def _x_options(at):
    return list([s for s in at.selectbox if s.label == "X"][0].options[1:])


class TestSwitchingSheets:
    """Changing the sheet must change the columns every control offers.

    Honest limitation: these tests pass with or without the sheet-scoped widget
    keys the page now uses. AppTest re-runs the script in-process and does not
    model the browser's widget-identity cache, which is the thing that can carry
    one sheet's selection onto another sheet's data when two sheets have
    identical column sets. So this class guards the *logic* -- reset, reload,
    re-offer -- and not the browser behaviour. Verified by neutralising the keys
    and watching all four still pass; do not read them as proof of that fix.
    """

    def test_columns_follow_the_selected_sheet(self, two_sheet_book):
        at = _run(two_sheet_book, sheet=None, header_rows=1)
        assert _x_options(at) == ["Alpha", "Beta"]

        _sheet_box(at).select("SheetB").run()
        assert not at.exception
        assert at.session_state["plot_explorer"]["sheet"] == "SheetB"
        assert _x_options(at) == ["Gamma", "Delta", "Epsilon"]

    def test_switching_back_restores_the_first_sheets_columns(self, two_sheet_book):
        at = _run(two_sheet_book, sheet=None, header_rows=1)
        _sheet_box(at).select("SheetB").run()
        _sheet_box(at).select("SheetA").run()
        assert not at.exception
        assert _x_options(at) == ["Alpha", "Beta"]

    def test_a_chosen_axis_does_not_survive_onto_a_different_sheet(self, two_sheet_book):
        at = _run(two_sheet_book, sheet=None, header_rows=1)
        [s for s in at.selectbox if s.label == "X"][0].select("Alpha").run()
        assert at.session_state["plot_explorer"]["config"]["x"] == "Alpha"

        _sheet_box(at).select("SheetB").run()
        assert at.session_state["plot_explorer"]["config"]["x"] is None

    def test_identical_columns_still_mean_a_separate_widget(self, two_sheet_book):
        """The collision case: same columns, different sheet, different data.

        With keyless widgets these two sheets hash to one widget and SheetA's
        axis silently reappears over SheetATwin's rows.
        """
        at = _run(two_sheet_book, sheet=None, header_rows=1)
        [s for s in at.selectbox if s.label == "X"][0].select("Alpha").run()

        _sheet_box(at).select("SheetATwin").run()
        assert not at.exception
        assert _x_options(at) == ["Alpha", "Beta"]
        assert at.session_state["plot_explorer"]["config"]["x"] is None


class TestPersistence:
    def test_the_config_is_remembered_for_the_sheet(self, workbook, tmp_path):
        at = _run(workbook, config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)"})
        assert not at.exception

        remembered = config_store.load_sheet_state("HA_SplitTable", tmp_path / "plot_explorer.json")
        assert remembered is not None
        assert remembered["config"]["x"] == "T (degC)"
        assert remembered["header_rows"] == 2

    def test_a_remembered_column_the_sheet_lacks_is_dropped(self, workbook):
        """A stale preference must not be able to break the page."""
        at = _run(workbook, config={**DEFAULT_CONFIG, "x": "T (degC)", "y": "FWHM (E2g)",
                                    "color": "a column from another sheet"})
        assert not at.exception
        assert not at.error
        assert len(at.get("plotly_chart")) == 1
